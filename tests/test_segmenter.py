"""统一结构读取分段器测试（一次加载与结构路由）。"""

import glob
import os
from pathlib import Path

from question_types.segmenter import (
    load_document_once,
    parse_document_once,
)

DOC_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "examples", "documents",
)


def _all_documents():
    return sorted(
        glob.glob(os.path.join(DOC_DIR, "*.docx"))
        + glob.glob(os.path.join(DOC_DIR, "*.xlsx"))
    )


def test_load_document_once_returns_paras_and_metadata():
    paras, metadata = load_document_once(
        os.path.join(DOC_DIR, "7上-U2-信息获取.docx"))
    assert len(paras) > 0
    assert len(paras) == len(metadata)


def test_parse_document_once_is_deterministic_for_all_samples():
    """所有示例文档都走唯一的新解析入口，并保持结果稳定。"""
    for path in _all_documents():
        first_results, first_summary = parse_document_once(path)
        second_results, second_summary = parse_document_once(path)
        assert first_results == second_results, path
        assert first_summary == second_summary, path


def test_unrecognized_document_reports_no_type(tmp_path):
    """未识别样例在测试内构造，避免依赖可被清理的示例二进制文件。"""
    from docx import Document

    path = tmp_path / "unrecognized.docx"
    document = Document()
    document.add_paragraph("This is an ordinary note without a supported question type.")
    document.save(path)

    results, summary = parse_document_once(path)
    assert results == []
    assert summary == "未识别到任何题型内容"


def test_preloaded_parser_reuses_document_load():
    """预加载段落注入后解析器不再自行加载（计数可验证只读一次）。"""
    calls = []
    original = load_document_once.__globals__["load_paragraphs"]

    def counting_load(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)

    load_document_once.__globals__["load_paragraphs"] = counting_load
    try:
        path = os.path.join(DOC_DIR, "S1-听后应答.docx")
        results, _ = parse_document_once(path)
        assert results and results[0]["item_count"] == 7
        assert len(calls) == 1, "文档应只加载一次"
    finally:
        load_document_once.__globals__["load_paragraphs"] = original


def test_xlsx_vocabulary_keeps_sheet_and_row_source_locations():
    """Excel 词汇条目进入核对页前，必须保留可定位的工作表/行号。"""
    path = os.path.join(DOC_DIR, "U6单词导入模板.xlsx")
    results, _ = parse_document_once(path)
    assert results and results[0]["doc_type"] == "词汇"
    first_word, first_sentence = results[0]["items"][:2]
    assert first_word["category"] == "单词"
    assert first_sentence["category"] == "例句"
    assert first_word["source_locator"].startswith("工作表/")
    assert "/行/2/单词" in first_word["source_locator"]
    assert "/行/2/例句" in first_sentence["source_locator"]


def test_xlsx_vocabulary_scans_matching_sheets_and_prefers_exact_word_header(tmp_path):
    """多工作表模板不能只读活动页，也不能把「单词释义」误当单词列。"""
    from openpyxl import Workbook

    path = tmp_path / "multi-sheet-vocabulary.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "第一单元"
    first.append(["单词释义", "单词名称", "例句"])
    first.append(["错误列", "pigeon", "A pigeon is here."])
    second = workbook.create_sheet("第二单元")
    second.append(["例句", "单词"])
    second.append(["A bird is there.", "bird"])
    workbook.save(path)
    workbook.close()

    results, summary = parse_document_once(str(path))
    items = results[0]["items"]
    assert summary == "检测到 1 种题型，成功提取 4 条内容"
    assert [item["text"] for item in items] == [
        "pigeon", "A pigeon is here.", "bird", "A bird is there.",
    ]
    assert "/第一单元/行/2/单词" in items[0]["source_locator"]
    assert "/第二单元/行/2/单词" in items[2]["source_locator"]


def test_parse_document_once_accepts_pathlike_xlsx():
    results, summary = parse_document_once(Path(DOC_DIR) / "U6单词导入模板.xlsx")
    assert results and results[0]["doc_type"] == "词汇"
    assert "成功提取" in summary
