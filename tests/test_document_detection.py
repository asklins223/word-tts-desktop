"""结构证据判型回归：文件名不参与题型路由。"""

from pathlib import Path
from shutil import copyfile

from docx import Document
from openpyxl import Workbook

import question_types
import workflow.parser as workflow_parser
from question_types import detect_document_type
from question_types.detection import (
    DocumentTypeEvidence,
    classify_paper_category,
    detect_document_types,
)
from question_types.segmenter import parse_document_once
from question_types.text_utils import load_paragraphs


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "examples/documents"


def test_real_path_uses_structure_even_when_name_is_neutral(tmp_path):
    source = DOC_DIR / "听后选择-7上 Starter Unit 1 Hello!.docx"
    target = tmp_path / "materials-001.docx"
    copyfile(source, target)

    evidence = detect_document_types(target)
    assert [item.doc_type for item in evidence] == ["听后选择"]
    assert evidence[0].question_count == 8
    assert evidence[0].option_count == 24
    assert detect_document_type(target) == "听后选择"

    results, summary = parse_document_once(target)
    assert summary == "检测到 1 种题型，成功提取 6 条内容"
    assert results[0]["doc_type"] == "听后选择"


def test_filename_markers_do_not_classify_an_unrelated_word_document(tmp_path):
    target = tmp_path / "听后选择-模仿朗读-套卷.docx"
    document = Document()
    document.add_paragraph("This is an ordinary note.")
    document.add_paragraph("The note mentions 听后选择 and 模仿朗读 only as words.")
    document.save(target)

    assert detect_document_types(target) == ()
    assert parse_document_once(target) == ([], "未识别到任何题型内容")


def test_imitation_heading_and_english_note_are_not_enough_for_a_type(tmp_path):
    target = tmp_path / "ordinary-note.docx"
    document = Document()
    document.add_paragraph("模仿朗读")
    document.add_paragraph("This is an ordinary note about a classroom activity.")
    document.save(target)

    assert detect_document_types(target) == ()


def test_content_detection_requires_structural_evidence():
    paragraphs = [(0, "一、听后选择", "Normal")]

    assert detect_document_types(paragraphs=paragraphs) == ()


def test_mixed_paper_reports_question_and_table_evidence_without_filename(tmp_path):
    source = DOC_DIR / "七上Starter Unit 1 听说测试题（2026新题型）.docx"
    target = tmp_path / "paper-without-a-type-name.docx"
    copyfile(source, target)

    evidence = detect_document_types(target)
    by_type = {item.doc_type: item for item in evidence}
    assert [item.doc_type for item in evidence] == [
        "听后选择",
        "听后应答",
        "模仿朗读",
        "听后记录并转述信息",
    ]
    assert by_type["听后选择"].question_count == 8
    assert by_type["听后选择"].option_count == 24
    assert "应答控制提示: 7" in by_type["听后应答"].signals
    assert by_type["模仿朗读"].question_count == 1
    assert by_type["听后记录并转述信息"].table_count == 2
    assert by_type["听后记录并转述信息"].details["recording_table_count"] == 1
    # A mixed document has no single owner, even when its basename contains no
    # useful hint.
    assert detect_document_type(target) is None


def test_paper_category_is_derived_from_document_structure_and_propagated_to_items():
    full_paper = DOC_DIR / "七上Starter Unit 1 听说测试题（2026新题型）.docx"
    full_paragraphs = load_paragraphs(full_paper)
    full_decision = classify_paper_category(full_paragraphs)

    assert full_decision.paper_category == "听说考试"
    assert full_decision.exam_form == "paper"
    assert full_decision.status == "suggested"
    assert {"听后选择", "听后应答"}.issubset(set(full_decision.detected_types))

    results, _ = parse_document_once(full_paper)
    parsed_items = [item for result in results for item in result["items"]]
    assert parsed_items
    assert {item["paper_category"] for item in parsed_items} == {"听说考试"}
    assert {item["exam_form"] for item in parsed_items} == {"paper"}
    assert {item["paper_category_status"] for item in parsed_items} == {"suggested"}

    normalized = workflow_parser.DocumentParser().parse(full_paper)
    assert {item.metadata["paper_category"] for item in normalized.items} == {"听说考试"}
    assert {item.metadata["exam_form"] for item in normalized.items} == {"paper"}
    imitation_items = [
        item for item in normalized.items
        if item.metadata.get("doc_type") == "模仿朗读"
    ]
    assert imitation_items
    assert {item.metadata["major_section_profile"] for item in imitation_items} == {
        "imitation_boxed_special"
    }
    assert {item.metadata["entry_profile"] for item in imitation_items} == {
        "imitation_reading_v1"
    }
    assert all(item.metadata["capabilities"]["external_input"] for item in imitation_items)

    special = DOC_DIR / "听后选择-7上 Starter Unit 1 Hello!.docx"
    special_decision = classify_paper_category(load_paragraphs(special))
    assert special_decision.paper_category == "题型专项"
    assert special_decision.exam_form == "special"
    assert special_decision.status == "suggested"

    special_results, _ = parse_document_once(special)
    special_items = [item for result in special_results for item in result["items"]]
    assert {item["paper_category"] for item in special_items} == {"题型专项"}
    assert {item["exam_form"] for item in special_items} == {"special"}

    imitation_special = workflow_parser.DocumentParser().parse(
        DOC_DIR / "七上Starter Unit1 Hello模仿朗读专项.docx"
    )
    assert imitation_special.items
    assert all(
        item.metadata.get("page_input", {}).get("type") == "模仿朗读"
        and item.metadata.get("entry_profile") == "imitation_reading_v1"
        and item.metadata.get("capabilities", {}).get("external_input") is True
        for item in imitation_special.items
    )


def test_multiple_types_without_full_paper_structure_remain_unresolved():
    decision = classify_paper_category(
        ["一份专项资料", "听后选择", "模仿朗读"],
        evidences=(
            DocumentTypeEvidence(
                code="listening_choice",
                doc_type="听后选择",
                score=8,
                structural=True,
            ),
            DocumentTypeEvidence(
                code="imitation_reading",
                doc_type="模仿朗读",
                score=8,
                structural=True,
            ),
        ),
    )

    assert decision.paper_category is None
    assert decision.exam_form == "unknown"
    assert decision.status == "conflict"


def test_xlsx_type_requires_word_and_sentence_headers(tmp_path):
    target = tmp_path / "单词导入模板-but-wrong.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["单词释义", "例句"])
    worksheet.append(["definition", "A sentence."])
    workbook.save(target)
    workbook.close()

    assert detect_document_types(target) == ()
    assert detect_document_type(target) is None


def test_missing_path_with_a_type_name_is_not_a_document_type(tmp_path):
    missing = tmp_path / "听后选择-not-a-real-document.docx"
    assert detect_document_type(missing) is None


def test_public_parser_surface_exposes_only_new_entrypoints():
    """Removed compatibility names must not leak back into the new package."""

    assert not hasattr(question_types, "detect_doc_type")
    assert not hasattr(question_types, "detect_types_in_content")
    assert not hasattr(question_types, "parse_document_auto")
    assert not hasattr(workflow_parser, "LegacyWordParser")
    assert hasattr(workflow_parser, "DocumentParser")


def test_record_table_questions_are_evidence_even_when_blanks_use_underscores():
    source = DOC_DIR / (
        "七上Starter Unit 1 Hello听后记录并转述信息专项-"
        "答案扩展(1).docx"
    )

    evidence = detect_document_types(source)
    assert evidence[0].question_count == 3
    assert evidence[0].details["table_question_count"] == 3


def test_auto_numbered_word_questions_are_part_of_structure_evidence():
    """Word list numbering must not hide the last questions from detection."""

    source = DOC_DIR / "7上-U2-信息获取.docx"
    paragraphs, paragraph_metadata, blocks = load_paragraphs(
        source,
        include_metadata=True,
        include_blocks=True,
    )

    evidence = detect_document_types(
        paragraphs=paragraphs,
        paragraph_metadata=paragraph_metadata,
        blocks=blocks,
    )

    assert [item.doc_type for item in evidence] == ["信息获取"]
    assert evidence[0].question_count == 10
    assert evidence[0].details["auto_numbered_question_count"] == 3
    assert evidence[0].details["question_numbers"] == list(range(1, 11))


def test_repeated_question_numbers_are_counted_as_separate_structural_items():
    """Repeated numbering is valid when a document contains repeated blocks."""

    paragraphs = [
        (0, "一、听后选择", "Heading 1"),
        (1, "1. What is it?", "Normal"),
        (2, "A. One.", "Normal"),
        (3, "B. Two.", "Normal"),
        (4, "录音原文", "Normal"),
        (5, "W: It is two.", "Normal"),
        (6, "一、听后选择", "Heading 1"),
        (7, "1. Where is it?", "Normal"),
        (8, "A. Here.", "Normal"),
        (9, "B. There.", "Normal"),
        (10, "录音原文", "Normal"),
        (11, "W: It is there.", "Normal"),
    ]

    evidence = detect_document_types(paragraphs=paragraphs)

    assert len(evidence) == 1
    assert evidence[0].question_count == 2
    assert evidence[0].details["question_numbers"] == [1, 1]
