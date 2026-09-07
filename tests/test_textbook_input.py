"""课文（textbook）系统录入的专项测试。

覆盖三块新增能力：解析器保留中文译文、文件名/结构驱动的配置建议、
课文页面执行器的本地 spec 构造校验。
"""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from workflow.artifact_store import ArtifactStore
from workflow.database import WorkflowDatabase
from workflow.repositories import WorkflowRepository
from workflow.system_input import SystemInputError, SystemInputService
from workflow.system_input_executor import TextbookInputWorkflowPageExecutor
from workflow.textbook_suggestions import suggest_textbook_configuration

from question_types import TextReadingParser


class TextbookSuggestionTests(unittest.TestCase):
    def test_filename_driven_classification_defaults(self) -> None:
        suggestions = suggest_textbook_configuration(
            "课文跟读-7上-Unit1 You and Me.docx",
            [],
        )
        self.assertEqual(suggestions["textbookGrade"]["value"], "七年级")
        self.assertEqual(suggestions["textbookStage"]["value"], "初中")
        self.assertEqual(suggestions["textbookVolume"]["value"], "上册")
        self.assertEqual(suggestions["textbookUnit"]["value"], "Unit 1")
        self.assertEqual(suggestions["textbookNameEn"]["value"], "You and Me")
        self.assertNotIn("textbookVersion", suggestions)

    def test_version_marker_and_g_style_grade(self) -> None:
        suggestions = suggest_textbook_configuration(
            "课文跟读人教九上U1.docx",
            [],
        )
        self.assertEqual(suggestions["textbookVersion"]["value"], "人教版")
        self.assertEqual(suggestions["textbookGrade"]["value"], "九年级")
        self.assertEqual(suggestions["textbookUnit"]["value"], "Unit 1")

    def test_tight_filename_keeps_volume(self) -> None:
        # “9上U1”“8下Unit3”这类册别紧跟年级、后面直接连字母的写法，
        # 册别不能因为 lookahead 被字母挡住而丢失。
        for filename, expected_volume in (
            ("课文跟读-9上U1.docx", "上册"),
            ("课文跟读-8下Unit3.docx", "下册"),
            ("课文跟读7上U2.docx", "上册"),
        ):
            suggestions = suggest_textbook_configuration(filename, [])
            self.assertEqual(
                suggestions["textbookVolume"]["value"],
                expected_volume,
                filename,
            )
            self.assertNotIn("textbookNameEn", suggestions)

    def test_ambiguous_lesson_keeps_empty_with_candidates(self) -> None:
        rows = [
            {"section": "Section A", "role": "Meime"},
            {"section": "Section B"},
        ]
        suggestions = suggest_textbook_configuration("课文跟读-7上.docx", rows)
        lesson = suggestions["textbookLesson"]
        self.assertEqual(lesson["value"], "")
        self.assertEqual(lesson["candidates"], ["Section A", "Section B"])
        self.assertEqual(suggestions["textbookForm"]["value"], "角色扮演")

    def test_role_metadata_drives_roleplay_suggestion(self) -> None:
        # role 是解析条目的独立列；normalize_item 必须把它透传进
        # metadata，课文形式建议才能区分角色扮演与同步课文。
        from workflow.parser import normalize_item

        item = normalize_item(
            {"text": "Teng Fei: Hello!", "category": "段落跟读", "role": "Teng Fei"},
            sequence=1,
            source_basis="basis",
            document_type="课文跟读",
        )
        self.assertEqual(item.metadata.get("role"), "Teng Fei")

        suggestions = suggest_textbook_configuration(
            "课文跟读-7上.docx",
            [item.metadata],
        )
        self.assertEqual(suggestions["textbookForm"]["value"], "角色扮演")

        plain = normalize_item(
            {"text": "Where are you from?", "category": "句子跟读"},
            sequence=2,
            source_basis="basis",
            document_type="课文跟读",
        )
        suggestions = suggest_textbook_configuration(
            "课文跟读-7上.docx",
            [plain.metadata],
        )
        self.assertEqual(suggestions["textbookForm"]["value"], "同步课文")

    def test_single_section_suggests_lesson(self) -> None:
        rows = [{"section": "Reading Plus"}, {"section": "Reading Plus"}]
        suggestions = suggest_textbook_configuration("课文跟读-9上-U1.docx", rows)
        self.assertEqual(suggestions["textbookLesson"]["value"], "Reading Plus")
        self.assertEqual(suggestions["textbookForm"]["value"], "同步课文")


class TextTranslationAttachmentTests(unittest.TestCase):
    def test_sentences_keep_chinese_translation_lines(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "课文跟读.docx"
            document = Document()
            for text in [
                "Section A",
                "句子跟读",
                "1. May I have your name?",
                "中文：请问你叫什么名字？",
                "2. Where are you from?",
                "中文：你来自哪里？",
            ]:
                document.add_paragraph(text)
            document.save(path)
            result = TextReadingParser(str(path)).parse()
        translations = [item.get("translation") for item in result["items"]]
        self.assertEqual(translations, ["请问你叫什么名字？", "你来自哪里？"])

    def test_discourse_does_not_inherit_sentence_translations(self) -> None:
        sentence = "Over the years, he collected many important seeds for China's seed banks."
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "课文跟读.docx"
            document = Document()
            for text in [
                "Section B",
                "句子跟读",
                f"1. {sentence}",
                "中文：多年来，他为中国的种子库收集了许多重要的种子。",
                "语篇跟读",
                "The Inspiring Seed Scientist",
                f"{sentence} He kept working in Xizang.",
            ]:
                document.add_paragraph(text)
            document.save(path)
            result = TextReadingParser(str(path)).parse()
        translations_by_category: dict[str, list] = {}
        for item in result["items"]:
            translations_by_category.setdefault(item["category"], []).append(item.get("translation"))
        self.assertEqual(
            translations_by_category["句子跟读"],
            ["多年来，他为中国的种子库收集了许多重要的种子。"],
        )
        self.assertEqual(
            translations_by_category["语篇跟读"],
            [None, None],
        )

    def test_discourse_keeps_its_own_translation_lines(self) -> None:
        sentence = "Over the years, he collected many important seeds for China's seed banks."
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "课文跟读.docx"
            document = Document()
            for text in [
                "Section B",
                "句子跟读",
                f"1. {sentence}",
                "中文：句子译文。",
                "语篇跟读",
                f"{sentence}",
                "中文：语篇译文。",
            ]:
                document.add_paragraph(text)
            document.save(path)
            result = TextReadingParser(str(path)).parse()
        translations_by_category: dict[str, list] = {}
        for item in result["items"]:
            translations_by_category.setdefault(item["category"], []).append(item.get("translation"))
        self.assertEqual(translations_by_category["句子跟读"], ["句子译文。"])
        self.assertEqual(translations_by_category["语篇跟读"], ["语篇译文。"])


class TextbookExecutorSpecTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="wordtts-textbook-input-")
        root = Path(self.temp.name)
        self.database = WorkflowDatabase(root / "workflow.db", profile="full")
        self.database.initialize()
        self.artifacts = ArtifactStore(root / "artifacts")
        self.repository = WorkflowRepository(self.database)
        self.executor = TextbookInputWorkflowPageExecutor(
            self.repository,
            self.artifacts,
            profile_dir=root / "profile",
            admin_url="https://admin.example.test",
            api_base="https://api.example.test",
            login_timeout=5,
        )
        self.workflow = self.repository.create_workflow(
            "tts",
            {
                "source_filename": "课文跟读-7上-Unit1 You and Me.docx",
                "delivery_mode": "audio_and_input",
                "input_type": "textbook",
            },
        )
        item_id = self.repository.create_item(
            self.workflow.workflow_id,
            item_type="text_reading",
            sequence=0,
            normalized_content="May I have your name?",
            item_identity_key="text:1",
            metadata={"doc_type": "课文跟读", "translation": "请问你叫什么名字？"},
            status="SUCCEEDED",
        )
        staged = self.artifacts.stage_stream(io.BytesIO(b"ID3-textbook-audio"))
        blob = self.artifacts.promote(staged, format="mp3")
        self.repository.attach_imported_artifact(
            self.workflow.workflow_id,
            artifact_id="textbook-audio-1",
            blob=blob,
            artifact_type="tts-segment",
            producer="test",
            producer_version="1",
            item_id=item_id,
        )
        self.item_id = item_id

    def tearDown(self) -> None:
        self.database.close()
        self.temp.cleanup()

    def _payload(self, item_id: str, *, translation: str = "请问你叫什么名字？") -> dict:
        return {
            "workflow_id": self.workflow.workflow_id,
            "unit": {
                "unit_id": "unit-1",
                "input_type": "textbook",
                "configuration": {
                    "input_type": "textbook",
                    "textbookNameZh": "Section A",
                    "textbookNameEn": "How do we get to know each other?",
                    "textbookForm": "角色扮演",
                    "textbookVersion": "人教版",
                    "textbookStage": "初中",
                    "textbookGrade": "七年级",
                    "textbookVolume": "上册",
                    "textbookUnit": "Unit 1",
                    "textbookLesson": "Section A",
                },
                "segments": [{
                    "segment_id": "seg-1",
                    "item_id": item_id,
                    "ordinal": 0,
                    "raw_text": "May I have your name?",
                    "tts_text": "May I have your name?",
                    "score": None,
                    "audio_artifact_id": "textbook-audio-1",
                    "translation": translation,
                }],
            },
            "target": {"workflow_id": self.workflow.workflow_id, "input_type": "textbook"},
        }

    def test_build_records_maps_configuration_and_segments(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            records = self.executor._build_records(
                self._payload(self.item_id),
                Path(temp),
            )
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["name_zh"], "Section A")
        self.assertEqual(record["name_en"], "How do we get to know each other?")
        self.assertEqual(record["form"], "角色扮演")
        self.assertEqual(record["grade"], "七年级")
        self.assertEqual(len(record["items"]), 1)
        item = record["items"][0]
        self.assertEqual(item["original"], "May I have your name?")
        self.assertEqual(item["translation"], "请问你叫什么名字？")
        self.assertTrue(item["audio_path"].endswith(".mp3"))

    def test_missing_classification_fields_fail_closed_locally(self) -> None:
        payload = self._payload(self.item_id)
        del payload["unit"]["configuration"]["textbookNameZh"]
        del payload["unit"]["configuration"]["textbookLesson"]
        with tempfile.TemporaryDirectory() as temp, self.assertRaises(SystemInputError) as error:
            self.executor._build_records(payload, Path(temp))
        self.assertEqual(error.exception.code, "SYSTEM_INPUT_CONFIG_INCOMPLETE")
        self.assertEqual(error.exception.details["fields"], ["textbookNameZh", "textbookLesson"])

    def test_segment_without_audio_fails_before_browser(self) -> None:
        payload = self._payload(self.item_id)
        payload["unit"]["segments"][0]["audio_artifact_id"] = ""
        with tempfile.TemporaryDirectory() as temp, self.assertRaises(SystemInputError) as error:
            self.executor._build_records(payload, Path(temp))
        self.assertEqual(error.exception.code, "AUDIO_GATE_FAILED")

    def test_paper_payload_is_rejected(self) -> None:
        payload = self._payload(self.item_id)
        payload["unit"]["input_type"] = "paper"
        payload["unit"]["configuration"]["input_type"] = "paper"
        with tempfile.TemporaryDirectory() as temp, self.assertRaises(SystemInputError) as error:
            self.executor._build_records(payload, Path(temp))
        self.assertEqual(error.exception.code, "SYSTEM_INPUT_TYPE_UNSUPPORTED")


class TextbookProjectionFlowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="wordtts-textbook-flow-")
        root = Path(self.temp.name)
        self.database = WorkflowDatabase(root / "workflow.db", profile="full")
        self.database.initialize()
        self.artifacts = ArtifactStore(root / "artifacts")
        self.repository = WorkflowRepository(self.database)
        self.external_service = __import__(
            "workflow.external", fromlist=["ExternalRecordService"]
        ).ExternalRecordService(self.database, intent_log=self.repository.intent_log)
        self.service = SystemInputService(
            self.database,
            self.repository,
            external=self.external_service,
        )
        # 注意：workflow 初始配置不写 input_type，让解析分类保持
        # suggested（未被用户保存覆盖）。
        self.workflow = self.repository.create_workflow(
            "tts",
            {
                "source_filename": "课文跟读-7上-Unit1 You and Me.docx",
                "delivery_mode": "audio_and_input",
            },
        )
        for sequence, (english, chinese) in enumerate([
            ("May I have your name?", "请问你叫什么名字？"),
            ("Where are you from?", "你来自哪里？"),
        ]):
            item_id = self.repository.create_item(
                self.workflow.workflow_id,
                item_type="text_reading",
                sequence=sequence,
                normalized_content=english,
                item_identity_key=f"text:{sequence + 1}",
                metadata={
                    "doc_type": "课文跟读",
                    "category": "句子跟读",
                    "section": "Section A",
                    "raw_text": english,
                    "translation": chinese,
                },
                status="SUCCEEDED",
            )
            staged = self.artifacts.stage_stream(io.BytesIO(f"audio-{sequence}".encode()))
            blob = self.artifacts.promote(staged, format="mp3")
            self.repository.attach_imported_artifact(
                self.workflow.workflow_id,
                artifact_id=f"textbook-audio-{sequence + 1}",
                blob=blob,
                artifact_type="tts-segment",
                producer="test",
                producer_version="1",
                item_id=item_id,
            )
        self.projection = self.service.sync_projection(
            self.workflow.workflow_id,
            source_filename="课文跟读-7上-Unit1 You and Me.docx",
        )

    def tearDown(self) -> None:
        self.database.close()
        self.temp.cleanup()

    def test_textbook_projection_exposes_supported_gate_and_suggestions(self) -> None:
        self.assertEqual(self.projection["input_type"], "textbook")
        self.assertEqual(self.projection["input_type_status"], "suggested")
        self.assertTrue(self.projection["document_entry_support"]["supported"])
        self.assertEqual(
            self.projection["document_entry_support"]["format"],
            "text_reading",
        )
        suggested = self.projection["suggested_configuration"]
        self.assertEqual(suggested["input_type"], "textbook")
        self.assertEqual(suggested["input_type_status"], "suggested")
        self.assertEqual(suggested["textbook"]["textbookGrade"]["value"], "七年级")
        self.assertEqual(suggested["textbook"]["textbookUnit"]["value"], "Unit 1")

    def test_saved_textbook_configuration_survives_projection(self) -> None:
        unit_id = self.projection["units"][0]["unit_id"]
        saved = self.service.save_configuration(
            self.workflow.workflow_id,
            self.repository.get_workflow(self.workflow.workflow_id).state_version,
            {
                "delivery_mode": "audio_and_input",
                "input_type": "textbook",
                "units": [{
                    "unit_id": unit_id,
                    "textbookNameZh": "Section A",
                    "textbookNameEn": "How do we get to know each other?",
                    "textbookForm": "角色扮演",
                    "textbookVersion": "人教版",
                    "textbookStage": "初中",
                    "textbookGrade": "七年级",
                    "textbookVolume": "上册",
                    "textbookUnit": "Unit 1",
                    "textbookLesson": "Section A",
                }],
            },
        )
        unit = saved["projection"]["units"][0]
        self.assertEqual(unit["input_type_status"], "user_override")
        self.assertEqual(unit["configuration"]["textbookNameZh"], "Section A")
        self.assertEqual(unit["configuration"]["textbookLesson"], "Section A")
        entry = saved["projection"]["entries"][0]
        self.assertEqual(entry["document_name"], "Section A")
        self.assertIsNone(entry["review_url"])
        self.assertEqual(entry["review_url_status"], "unknown")
        # 段落片段携带译文，音频闸门产物已映射到片段。
        for segment in saved["projection"]["content_segments"]:
            self.assertTrue(segment.get("translation"))
            self.assertTrue(segment.get("audio_artifact_id"))

    def _finish_audio(self, workflow_id: str) -> None:
        with self.database.transaction() as con:
            con.execute(
                """UPDATE workflows SET status='ACTIVE', execution_state='TERMINAL',
                   control_state='TERMINATED', cleanup_state='SUCCEEDED',
                   result_status='SUCCEEDED', finished_at=updated_at
                   WHERE workflow_id=?""",
                (workflow_id,),
            )

    def test_full_textbook_run_delivers_translation_and_records_success(self) -> None:
        unit_id = self.projection["units"][0]["unit_id"]
        saved = self.service.save_configuration(
            self.workflow.workflow_id,
            self.repository.get_workflow(self.workflow.workflow_id).state_version,
            {
                "delivery_mode": "audio_and_input",
                "input_type": "textbook",
                "units": [{
                    "unit_id": unit_id,
                    "textbookNameZh": "Section A",
                    "textbookNameEn": "How do we get to know each other?",
                    "textbookForm": "同步课文",
                    "textbookVersion": "人教版",
                    "textbookStage": "初中",
                    "textbookGrade": "七年级",
                    "textbookVolume": "上册",
                    "textbookUnit": "Unit 1",
                    "textbookLesson": "Section A",
                }],
            },
        )
        self._finish_audio(self.workflow.workflow_id)
        accepted = self.service.accept_audio(
            self.workflow.workflow_id,
            saved["snapshot"]["state_version"],
        )
        started = self.service.start_input(
            self.workflow.workflow_id,
            accepted["snapshot"]["state_version"],
            idempotency_key="textbook-full-run-00000001",
        )

        payloads: list[dict] = []
        self.service.executor = lambda payload: payloads.append(payload) or {
            "status": "SUCCEEDED",
            "external_record_id": "tb-record-1",
        }
        result = self.service.execute_input_run(started["input_run_id"])

        self.assertEqual(result["input_status"], "succeeded")
        self.assertEqual(len(payloads), 1)
        payload = payloads[0]
        self.assertEqual(payload["target"]["input_type"], "textbook")
        unit = payload["unit"]
        self.assertEqual(unit["configuration"]["textbookNameZh"], "Section A")
        self.assertEqual(unit["configuration"]["textbookLesson"], "Section A")
        self.assertNotIn("paperName", unit["configuration"])
        self.assertEqual(len(unit["segments"]), 2)
        self.assertEqual(unit["segments"][0]["translation"], "请问你叫什么名字？")
        self.assertTrue(unit["segments"][0]["audio_artifact_id"])
        succeeded = [entry for entry in result["entries"] if entry["input_status"] == "succeeded"]
        self.assertEqual(len(succeeded), 1)
        self.assertEqual(succeeded[0]["external_record_id"], "tb-record-1")
        self.assertEqual(succeeded[0]["document_name"], "Section A")
        self.assertIsNone(succeeded[0]["review_url"])
        self.assertEqual(succeeded[0]["review_url_status"], "unavailable")
        self.assertEqual(succeeded[0]["review_url_source"], "platform_not_supported")


if __name__ == "__main__":
    unittest.main()
