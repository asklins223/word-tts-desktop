from pathlib import Path

from docx import Document
from PIL import Image

from question_types.info_retelling import InfoRetellingParser
from question_types.segmenter import parse_document_once
from question_types.text_utils import load_paragraphs
from workflow.parser import DocumentParser


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "examples/documents/信息转述及询问信息 7上- U1.docx"


def test_info_retelling_marks_the_positioned_drawing_group():
    results, summary = parse_document_once(FIXTURE)

    assert summary == "检测到 1 种题型，成功提取 1 条内容"
    assert len(results) == 1
    recording = results[0]["recording"]
    assert recording["block_image_required"] is True
    assert recording["block_kind"] == "drawing_group"
    assert recording["block_index"] == 0

def test_workflow_page_facts_keep_the_drawing_selector():
    parsed = DocumentParser().parse(FIXTURE)
    recordings = [item.metadata["page_input"]["recording"] for item in parsed.items]

    assert [recording["block_image_required"] for recording in recordings] == [True]
    assert [recording["block_kind"] for recording in recordings] == ["drawing_group"]
    assert [recording["block_index"] for recording in recordings] == [0]
    assert "(W)" not in recordings[0]["listening_text"]
    assert parsed.items[0].metadata["page_input"]["asking"][0]["reference_answers"] == [
        "What problem may you have with your plan?"
    ]
    recording = parsed.items[0].metadata["page_input"]["recording"]
    assert recording["asking_instruction_text"].startswith("你希望了解更多")
    assert recording["asking_instruction_occurrence"] == 0
    assert "asking_instruction_audio_filename_stem" not in recording


def test_retelling_sections_keep_distinct_group_indexes_before_deduplication():
    loaded = load_paragraphs(FIXTURE, include_metadata=True, include_blocks=True)
    parser = InfoRetellingParser(FIXTURE, preloaded_paras=loaded)

    assert parser._recording_drawing_group_indices() == (0, 1)


def test_retelling_selector_prefers_inline_image_in_its_subsection(tmp_path):
    source = tmp_path / "inline-retelling.docx"
    image_source = tmp_path / "mind-map.png"
    Image.new("RGB", (32, 24), "#2f7e69").save(image_source, format="PNG")
    document = Document()
    document.add_paragraph("第一节 信息转述")
    document.add_paragraph("请根据思维导图复述。")
    document.add_picture(str(image_source))
    document.add_paragraph("第二节 询问信息")
    document.save(source)

    loaded = load_paragraphs(source, include_metadata=True, include_blocks=True)
    parser = InfoRetellingParser(str(source), preloaded_paras=loaded)

    assert parser._recording_block_image_selector() == ("embedded_image", (0,))


def test_workflow_preserves_three_guide_audio_items_and_entry_facts(tmp_path):
    source = tmp_path / "three-guides.docx"
    document = Document()
    intro = "你将听到一篇 Peter 讲述去阳朔旅行的短文，录音播放两遍。"
    reading = "现在，你有15秒钟的时间阅读思维导图。"
    preparation = "下面你有60秒钟的时间准备，你的复述可以这样开始：Peter visited Yangshuo ..."
    for text in (
        "第一节 短文复述（共8分）", intro, reading, preparation,
        "听力原文：", "I'm Peter. I visited Yangshuo last August.",
        "参考答案：", "Peter visited Yangshuo last August.",
    ):
        document.add_paragraph(text)
    document.save(source)

    parsed = DocumentParser().parse(source, include_auxiliary_audio=True)
    guides = [item for item in parsed.items if item.metadata.get("audio_only_auxiliary")]
    assert [item.normalized_content for item in guides] == [intro, reading, preparation]
    assert len(parsed.items) == 4
    recording = parsed.items[0].metadata["page_input"]["recording"]
    assert recording["reading_instruction_text"] == reading
    assert recording["reading_instruction_audio_filename_stem"] == "信息转述阅图指导文字-1"
    assert recording["reading_instruction_occurrence"] == 1
    assert all(item.role == "题干音色" for item in guides)
