"""Content and structure based document/type detection.

The parser registry contains human-facing names and structural content markers;
a filename is never evidence that a document contains a type.  This module is
the single runtime boundary for document detection:

* Word documents are classified from headings, question/option layout,
  recording controls, English payload, and table structure;
* Excel vocabulary documents are classified from workbook headers, never from
  an ``.xlsx`` suffix alone;
* mixed exam papers may return more than one type, in document order;
* every public entry point uses a real source or already-loaded document
  structures; basename-only type guessing is not supported.

The detector deliberately does not import parser implementations.  That keeps
it usable before parser dispatch, avoids import cycles, and makes adding a new
parser possible through the generic registry fallback.  Existing families get
small structural profiles here because their source formats are materially
different (choice, response, reading, table-based recording, and so on).
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping

from audio_naming import is_exam_paper_bundle
from question_model.model import FAMILY_REGISTRY

from .text_utils import (
    is_chinese,
    is_major_section_heading,
    match_answer_marker,
    match_script_marker,
)


_NUMBERED_LINE_RE = re.compile(
    r"^\s*(?P<number>\d+)\s*[.．、）)]\s*(?P<body>.+?)\s*$"
)
_OPTION_LINE_RE = re.compile(
    r"^\s*[A-CＡ-Ｃ]\s*[.．、）)]\s*(?P<body>.+?)\s*$",
    re.IGNORECASE,
)
_STAR_LINE_RE = re.compile(r"^\s*[★☆*]")
_SPEAKER_RE = re.compile(r"^\s*(?:[WwMm]\s*[:：]|\([WwMm]\))")
_RESPONSE_PROMPT_RE = re.compile(
    r"听下面\s*(?:[0-9０-９零〇一二两三四五六七八九十百]+)?\s*个?\s*句子",
    re.IGNORECASE,
)
_ANSWER_PROMPT_RE = re.compile(r"请朗读应答语|朗读应答语", re.IGNORECASE)
_LISTENING_PROMPT_RE = re.compile(
    r"听下面|听第.+段|录音播放|各段播放|每段播放|短文听",
    re.IGNORECASE,
)
_IMITATION_CONTROL_RE = re.compile(
    r"听以下|听下面|准备|开始录音|停止录音|录音播放|计算机|模仿朗读",
    re.IGNORECASE,
)
_RECORDING_CONTROL_RE = re.compile(
    r"计算机|屏幕|答题区域|参考答案|答题时间|倒计时|"
    r"开始答题|停止转述|听短文|转述准备|完成转述",
    re.IGNORECASE,
)
_READING_ROLE_RE = re.compile(r"^([^:：\n]{1,60}?)\s*[:：]\s*.+$")
_READING_SUBSECTIONS = frozenset({"句子跟读", "段落跟读", "语篇跟读"})
_READING_CONTEXT_RE = re.compile(
    r"^(?:section\s+[a-z]|understanding\s+idea|reading\s+for\s+writing|"
    r"developing\s+ideas|reading\s+plus)\s*[：:]?$",
    re.IGNORECASE,
)
_SOURCE_LABEL_RE = re.compile(r"^\s*(?:外网|教材)\s*[：:]\s*$")
_RECORD_SECTION_RE = re.compile(
    r"第[一二三四五六七八九十百\d０-９]+节\s*[：:]?\s*听后记录"
)
_RECORD_TITLE_RE = re.compile(r"听后记录并转述信息")
_RETELLING_SECTION_RE = re.compile(
    r"第[一二三四五六七八九十百\d０-９]+节\s*[：:]?\s*信息转述"
)
_ASKING_SECTION_RE = re.compile(r"第[二2２]节\s*[：:]?\s*询问信息")
_INFO_ACQUISITION_HEADING_RE = re.compile(
    r"信息获取|听选信息|回答问题"
)
_SELECTION_HEADING_RE = re.compile(r"听后选择")
_RESPONSE_HEADING_RE = re.compile(r"听后应答")
_IMITATION_HEADING_RE = re.compile(r"模仿朗读")
_TEXT_TITLE_RE = re.compile(r"课文跟读")
_ASSESSMENT_LINE_RE = re.compile(
    r"评分标准|发音准确|停顿与语调|朗读流畅|错读|扣分|总分\s*\d+\s*分"
)


@dataclass(frozen=True)
class DocumentTypeEvidence:
    """Auditable evidence for one detected document family.

    Positions refer to the compact paragraph sequence returned by
    ``load_paragraphs``.  ``signals`` is intentionally human-readable so a
    UI or a future review report can explain why a family was selected rather
    than exposing an opaque filename guess.
    """

    code: str
    doc_type: str
    score: float
    start_positions: tuple[int, ...] = ()
    question_count: int = 0
    option_count: int = 0
    payload_count: int = 0
    table_count: int = 0
    structural: bool = False
    signals: tuple[str, ...] = ()
    details: Mapping[str, Any] = field(default_factory=dict, compare=False)

    @property
    def confidence(self) -> str:
        """Return a stable coarse confidence label for callers and logs."""

        if not self.structural:
            return "weak"
        if self.score >= 8:
            return "high"
        if self.score >= 5:
            return "medium"
        return "low"

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "doc_type": self.doc_type,
            "score": self.score,
            "confidence": self.confidence,
            "start_positions": list(self.start_positions),
            "question_count": self.question_count,
            "option_count": self.option_count,
            "payload_count": self.payload_count,
            "table_count": self.table_count,
            "structural": self.structural,
            "signals": list(self.signals),
            "details": dict(self.details),
        }


@dataclass(frozen=True)
class PaperCategoryEvidence:
    """Explain the document-level专项/套卷 classification.

    ``paper_category`` is the compatibility label consumed by the current
    system-input form.  ``exam_form`` is the parser-layer vocabulary from the
    document model: ``special`` means题型专项, ``paper`` means听说考试/套卷,
    and ``unknown`` means that the source does not provide enough independent
    structure to choose safely.
    """

    paper_category: str | None
    exam_form: str
    status: str
    confidence: str
    detected_types: tuple[str, ...] = ()
    signals: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return {
            "paper_category": self.paper_category,
            "exam_form": self.exam_form,
            "status": self.status,
            "confidence": self.confidence,
            "detected_types": list(self.detected_types),
            "signals": list(self.signals),
        }


@dataclass(frozen=True)
class _Corpus:
    paragraphs: tuple[tuple[int, str, str], ...]
    paragraph_metadata: tuple[Mapping[str, Any], ...] = ()
    blocks: tuple[Any, ...] = ()
    block_by_original_paragraph: Mapping[int, int] = field(default_factory=dict)
    major_positions: tuple[int, ...] = ()

    def metadata_at(self, position: int) -> Mapping[str, Any]:
        """Return compact-paragraph metadata without exposing list bounds."""

        if 0 <= position < len(self.paragraph_metadata):
            value = self.paragraph_metadata[position]
            return value if isinstance(value, Mapping) else {}
        return {}

    def window(self, starts: Iterable[int]) -> tuple[int, int]:
        """Return one family range, skipping repeated headings of that family."""

        values = sorted({int(value) for value in starts})
        if not values:
            return 0, 0
        start = values[0]
        family_starts = set(values)
        end = len(self.paragraphs)
        for position in self.major_positions:
            if position > start and position not in family_starts:
                end = position
                break
        return start, end

    def table_blocks(self, start: int, end: int) -> tuple[Any, ...]:
        """Return tables whose document-block positions fall in a paragraph range."""

        if not self.blocks:
            return ()
        if start >= len(self.paragraphs):
            return ()

        first_original = self.paragraphs[start][0]
        start_block = self.block_by_original_paragraph.get(first_original, 0)
        if end < len(self.paragraphs):
            last_original = self.paragraphs[end][0]
            end_block = self.block_by_original_paragraph.get(
                last_original, len(self.blocks)
            )
        else:
            end_block = len(self.blocks)
        return tuple(
            block
            for block in self.blocks
            if _block_kind(block) == "table"
            and start_block <= _block_index(block) < end_block
        )


@dataclass(frozen=True)
class _Features:
    question_numbers: tuple[int, ...] = ()
    auto_numbered_question_count: int = 0
    option_count: int = 0
    star_rows: int = 0
    script_count: int = 0
    answer_marker_count: int = 0
    response_prompt_count: int = 0
    answer_prompt_count: int = 0
    listening_prompt_count: int = 0
    control_count: int = 0
    speaker_count: int = 0
    english_count: int = 0
    score_signal_count: int = 0
    table_count: int = 0
    english_table_count: int = 0
    record_table_count: int = 0
    table_question_count: int = 0

    @property
    def question_count(self) -> int:
        return len(self.question_numbers)

    @property
    def payload_count(self) -> int:
        return (
            self.script_count
            + self.english_count
            + self.english_table_count
            + self.star_rows
        )


def _paragraphs_from(value: Iterable[Any]) -> tuple[tuple[int, str, str], ...]:
    result = []
    for position, paragraph in enumerate(value):
        if isinstance(paragraph, (tuple, list)):
            original = paragraph[0] if len(paragraph) > 0 else position
            text = paragraph[1] if len(paragraph) > 1 else ""
            style = paragraph[2] if len(paragraph) > 2 else ""
        else:
            original, text, style = position, paragraph, ""
        try:
            original = int(original)
        except (TypeError, ValueError):
            original = position
        text = str(text or "").strip()
        if text:
            result.append((original, text, str(style or "")))
    return tuple(result)


def _block_kind(block: Any) -> str:
    if isinstance(block, Mapping):
        return str(block.get("kind") or "")
    return str(getattr(block, "kind", "") or "")


def _block_index(block: Any) -> int:
    value = block.get("index", 0) if isinstance(block, Mapping) else getattr(block, "index", 0)
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _block_text(block: Any) -> str:
    if isinstance(block, Mapping):
        value = block.get("text", "")
    else:
        value = getattr(block, "text", "")
    return str(value or "").strip()


def _block_fragments(block: Any) -> tuple[str, ...]:
    value = (
        block.get("fragments", ())
        if isinstance(block, Mapping)
        else getattr(block, "fragments", ())
    )
    if isinstance(value, (list, tuple)):
        return tuple(str(item or "").strip() for item in value if str(item or "").strip())
    return ()


def _build_corpus(
    paragraphs: Iterable[Any],
    blocks: Iterable[Any] | None = None,
    paragraph_metadata: Iterable[Any] | None = None,
) -> _Corpus:
    normalized = _paragraphs_from(paragraphs)
    metadata_values = tuple(
        value if isinstance(value, Mapping) else {}
        for value in (paragraph_metadata or ())
    )
    # ``load_paragraphs`` returns metadata aligned to the non-empty paragraph
    # list.  Keep the detector tolerant of older/custom callers that provide
    # fewer entries, but never let an overlong metadata sequence shift later
    # paragraphs onto the wrong numbering facts.
    if len(metadata_values) < len(normalized):
        metadata_values += ({},) * (len(normalized) - len(metadata_values))
    elif len(metadata_values) > len(normalized):
        metadata_values = metadata_values[:len(normalized)]
    block_values = tuple(blocks or ())
    block_by_original = {}
    for block in block_values:
        if _block_kind(block) != "paragraph":
            continue
        metadata = (
            block.get("metadata", {})
            if isinstance(block, Mapping)
            else getattr(block, "metadata", {})
        )
        if not isinstance(metadata, Mapping):
            continue
        original = metadata.get("paragraph_index")
        try:
            original = int(original)
        except (TypeError, ValueError):
            continue
        block_by_original[original] = _block_index(block)

    # A Word style is useful evidence that a line is a heading, but it is not
    # by itself a cross-family boundary.  In reading documents every
    # ``Understanding Idea``/subsection line may be Heading 1/2/3; treating
    # those styles as global boundaries would truncate the reading payload
    # before the first sentence.  Only semantic major-section shapes split a
    # family window.
    major_positions = tuple(
        position
        for position, (_, text, _) in enumerate(normalized)
        if is_major_section_heading(text)
    )
    return _Corpus(
        paragraphs=normalized,
        paragraph_metadata=metadata_values,
        blocks=block_values,
        block_by_original_paragraph=block_by_original,
        major_positions=major_positions,
    )


def _is_heading_like(text: str, style: str = "") -> bool:
    value = str(text or "").strip()
    if not value:
        return False
    if is_major_section_heading(value):
        return True
    if re.search(r"heading|title|标题|副标题", str(style or ""), re.IGNORECASE):
        return True
    if value in _READING_SUBSECTIONS or _READING_CONTEXT_RE.match(value):
        return True
    return False


def _has_marker(patterns: Iterable[Any], value: str) -> bool:
    for pattern in patterns:
        try:
            if pattern.search(value):
                return True
        except (AttributeError, TypeError, re.error):
            if str(pattern) in value:
                return True
    return False


def _english_payload(value: str) -> bool:
    """Whether a line contains likely source text rather than UI copy."""

    value = str(value or "").strip()
    if len(re.findall(r"[A-Za-z]", value)) < 2:
        return False
    if _RECORDING_CONTROL_RE.search(value) or _IMITATION_CONTROL_RE.fullmatch(value):
        return False
    if match_script_marker(value) or match_answer_marker(value):
        return False
    # A mixed Chinese instruction with a few English abbreviations is not a
    # reading payload.  Long English clauses may legitimately contain a name
    # or a Chinese gloss, so only reject strongly Chinese lines here.
    if is_chinese(value) and len(re.findall(r"[A-Za-z]", value)) < 8:
        return False
    if _OPTION_LINE_RE.match(value):
        return False
    return True


def _table_features(corpus: _Corpus, start: int, end: int) -> tuple[int, int, int, int]:
    tables = corpus.table_blocks(start, end)
    english_tables = 0
    record_tables = 0
    table_question_numbers = set()
    for table in tables:
        text = _block_text(table)
        fragments = _block_fragments(table)
        if _english_payload(text) or len(re.findall(r"[A-Za-z]", text)) >= 8:
            english_tables += 1
        # Word tables often replace blank lines with the question number and
        # spaces, rather than literal underscores.  Bullets plus numbered
        # blanks are therefore stronger evidence than an underscore-only test.
        answer_key_only = bool(
            len(fragments) == 1
            and re.fullmatch(
                r"\s*(?:\d{1,3}\s*[.．、）)]?\s*_{2,}\s*){2,}",
                text,
            )
        )
        if (
            re.search(r"(?:_{2,}|\b\d{1,3}\b)", text)
            and not answer_key_only
            and (
                "●" in text
                or "·" in text
                or len(fragments) > 1
            )
        ):
            record_tables += 1
            table_question_numbers.update(
                int(value)
                for value in re.findall(r"(?<!\d)\d{1,3}(?!\d)", text)
                if int(value) > 0
            )
    return len(tables), english_tables, record_tables, len(table_question_numbers)


def _features(corpus: _Corpus, start: int, end: int) -> _Features:
    numbers = []
    auto_numbered_question_count = 0
    options = 0
    star_rows = 0
    scripts = 0
    answers = 0
    response_prompts = 0
    answer_prompts = 0
    listening_prompts = 0
    controls = 0
    speakers = 0
    english = 0
    score_signals = 0
    answer_region = False

    for position, (_, raw, _) in enumerate(corpus.paragraphs[start:end], start=start):
        value = str(raw or "").strip()
        if not value:
            continue
        script = match_script_marker(value)
        answer = match_answer_marker(value)
        if script:
            scripts += 1
            answer_region = False
        if answer:
            answers += 1
            answer_region = True
            continue
        if _RESPONSE_PROMPT_RE.search(value):
            response_prompts += 1
            listening_prompts += 1
        elif _LISTENING_PROMPT_RE.search(value):
            listening_prompts += 1
        if _ANSWER_PROMPT_RE.search(value):
            answer_prompts += 1
        if _RECORDING_CONTROL_RE.search(value) or _IMITATION_CONTROL_RE.search(value):
            controls += 1
        if _SPEAKER_RE.match(value):
            speakers += 1
        if re.search(r"每(?:小题|道题|题)|满分|分值|共\s*\d+\s*题", value):
            score_signals += 1
        if _STAR_LINE_RE.match(value):
            star_rows += 1
        option = _OPTION_LINE_RE.match(value)
        if option:
            options += 1
        numbered = _NUMBERED_LINE_RE.match(value)
        auto_numbered = False
        number = None
        body = value
        if numbered:
            body = numbered.group("body").strip()
            try:
                number = int(numbered.group("number"))
            except ValueError:
                number = None
        else:
            # Word's automatic list marker is not part of ``Paragraph.text``.
            # The shared loader preserves the rendered decimal number in
            # paragraph metadata; use it only for a question-like line, not
            # for an option row or a starred response row.
            raw_number = corpus.metadata_at(position).get("numbering_number")
            if (
                raw_number not in (None, "")
                and not _OPTION_LINE_RE.match(value)
                and not _STAR_LINE_RE.match(value)
            ):
                try:
                    number = int(raw_number)
                except (TypeError, ValueError):
                    number = None
                auto_numbered = number is not None

        if number is not None:
            likely_question = bool(re.search(r"[?？]\s*$", body))
            if (
                not body.startswith(("★", "☆", "*"))
                and not _RECORDING_CONTROL_RE.search(body)
                and not _ASSESSMENT_LINE_RE.search(body)
                and not _SPEAKER_RE.match(body)
                and (not answer_region or likely_question)
            ):
                # Automatic numbering is the parser's structural question
                # evidence even when a source omits the final question mark;
                # it is already constrained by the non-option/non-control
                # checks above.  Manual answer-region rows still require a
                # question mark so answer keys cannot become new questions.
                if not answer_region or likely_question:
                    numbers.append(number)
                    if auto_numbered:
                        auto_numbered_question_count += 1
        if not answer_region and _english_payload(value):
            english += 1

    (
        table_count,
        english_table_count,
        record_table_count,
        table_question_count,
    ) = _table_features(
        corpus, start, end
    )
    return _Features(
        question_numbers=tuple(numbers),
        auto_numbered_question_count=auto_numbered_question_count,
        option_count=options,
        star_rows=star_rows,
        script_count=scripts,
        answer_marker_count=answers,
        response_prompt_count=response_prompts,
        answer_prompt_count=answer_prompts,
        listening_prompt_count=listening_prompts,
        control_count=controls,
        speaker_count=speakers,
        english_count=english,
        score_signal_count=score_signals,
        table_count=table_count,
        english_table_count=english_table_count,
        record_table_count=record_table_count,
        table_question_count=table_question_count,
    )


def _positions(corpus: _Corpus, predicate) -> tuple[int, ...]:
    return tuple(
        position
        for position, (_, text, style) in enumerate(corpus.paragraphs)
        if predicate(text, style)
    )


def _heading_positions(corpus: _Corpus, pattern: Any) -> tuple[int, ...]:
    return _positions(
        corpus,
        lambda text, style: _is_heading_like(text, style)
        and _has_marker((pattern,), text),
    )


def _evidence(
    family,
    starts: Iterable[int],
    features: _Features,
    *,
    structural: bool,
    score: float,
    signals: Iterable[str],
    details: Mapping[str, Any] | None = None,
    question_count: int | None = None,
) -> DocumentTypeEvidence:
    feature_details = {
        "question_numbers": list(features.question_numbers),
        "auto_numbered_question_count": features.auto_numbered_question_count,
        "option_count": features.option_count,
        "star_rows": features.star_rows,
        "script_count": features.script_count,
        "answer_marker_count": features.answer_marker_count,
        "response_prompt_count": features.response_prompt_count,
        "answer_prompt_count": features.answer_prompt_count,
        "listening_prompt_count": features.listening_prompt_count,
        "control_count": features.control_count,
        "speaker_count": features.speaker_count,
        "english_count": features.english_count,
        "score_signal_count": features.score_signal_count,
        "table_count": features.table_count,
        "english_table_count": features.english_table_count,
        "record_table_count": features.record_table_count,
        "table_question_count": features.table_question_count,
    }
    feature_details.update(dict(details or {}))
    return DocumentTypeEvidence(
        code=family.code,
        doc_type=family.display_name,
        score=score,
        start_positions=tuple(sorted(set(starts))),
        question_count=(
            features.question_count
            if question_count is None
            else max(0, int(question_count))
        ),
        option_count=features.option_count,
        payload_count=features.payload_count,
        table_count=features.table_count,
        structural=structural,
        signals=tuple(dict.fromkeys(str(value) for value in signals)),
        details=feature_details,
    )


def _detect_info_acquisition(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _INFO_ACQUISITION_HEADING_RE)
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    has_section = any(
        re.search(r"听选信息|回答问题", corpus.paragraphs[position][1])
        for position in starts
    )
    structural = bool(
        has_section
        and features.question_count >= 1
        and (
            features.script_count >= 1
            or features.listening_prompt_count >= 1
        )
        and (
            features.option_count >= 2
            or features.answer_marker_count >= 1
            # The legacy format often uses slash-delimited options that do
            # not look like A-C rows.  A numbered question plus an explicit
            # speaker-marked script is still a complete acquisition block.
            or (features.script_count >= 1 and features.speaker_count >= 1)
        )
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=2 + features.question_count + features.script_count * 2,
        signals=(
            "题型标题/节标题: 信息获取或听选信息",
            f"题号结构: {features.question_count}",
            f"录音/听力提示: {features.script_count + features.listening_prompt_count}",
            f"选项/答案结构: {features.option_count + features.answer_marker_count}",
        ),
    )


def _detect_selection(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _SELECTION_HEADING_RE)
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    structural = bool(
        features.question_count >= 1
        and features.option_count >= 2
        and (
            features.script_count >= 1
            or features.listening_prompt_count >= 1
        )
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=3 + features.question_count + features.option_count / 3 + features.script_count * 2,
        signals=(
            "题型标题: 听后选择",
            f"编号题干: {features.question_count}",
            f"A-C 选项: {features.option_count}",
            f"录音稿/对话提示: {features.script_count + features.listening_prompt_count}",
        ),
    )


def _detect_response(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _RESPONSE_HEADING_RE)
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    structural = bool(
        features.response_prompt_count >= 1
        and features.answer_prompt_count >= 1
        and (features.star_rows >= 1 or features.english_count >= 1)
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=3 + features.response_prompt_count + features.answer_prompt_count + features.star_rows,
        signals=(
            "题型标题: 听后应答",
            f"听句提示: {features.response_prompt_count}",
            f"应答控制提示: {features.answer_prompt_count}",
            f"带星应答选项: {features.star_rows}",
        ),
        question_count=features.response_prompt_count or features.question_count,
    )


def _detect_text_reading(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _TEXT_TITLE_RE)
    starts = tuple(sorted(set(starts + _positions(
        corpus,
        lambda text, style: str(text).strip() in _READING_SUBSECTIONS,
    ))))
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    subsection_count = sum(
        str(text).strip() in _READING_SUBSECTIONS
        for _, text, _ in corpus.paragraphs[start:end]
    )
    context_count = sum(
        bool(_READING_CONTEXT_RE.match(str(text).strip()))
        for _, text, _ in corpus.paragraphs[start:end]
    )
    structural = bool(
        subsection_count >= 1
        and (
            features.english_count >= 1
            or features.english_table_count >= 1
            or features.speaker_count >= 2
        )
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=2 + subsection_count * 2 + features.english_count + context_count,
        signals=(
            f"跟读子题型标题: {subsection_count}",
            f"课文章节上下文: {context_count}",
            f"英文/角色载荷: {features.english_count + features.english_table_count}",
        ),
    )


def _record_context(corpus: _Corpus, position: int) -> bool:
    """Whether a paragraph is inside the new record/retelling paper section."""

    latest = None
    for candidate, (_, text, style) in enumerate(corpus.paragraphs[:position + 1]):
        value = str(text or "").strip()
        is_record_heading = _is_heading_like(value, style) and (
            _RECORD_TITLE_RE.search(value) or _RECORD_SECTION_RE.search(value)
        )
        if is_record_heading:
            latest = candidate
            continue
        if candidate == position and (
            _RETELLING_SECTION_RE.search(value)
            and re.search(r"第二节", value)
        ):
            # The second section is a child of the immediately preceding
            # record section, not an independent legacy retelling family.
            return latest is not None
        if candidate == position and is_major_section_heading(value):
            latest = None
    return latest is not None


def _detect_record_retelling(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _RECORD_TITLE_RE)
    starts = tuple(sorted(set(starts + _positions(
        corpus,
        lambda text, style: _is_heading_like(text, style)
        and bool(_RECORD_SECTION_RE.search(text)),
    ))))
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    record_instruction_count = sum(
        bool(_RECORDING_CONTROL_RE.search(text) or "信息记录表" in text or "一空限填" in text)
        for _, text, _ in corpus.paragraphs[start:end]
    )
    # The table is the normal source structure for this family.  A
    # paragraph-only import can still be meaningful, but then it must contain
    # more than one record-specific control plus an actual listening payload.
    # A title, one generic “参考答案/计算机” line, and an unrelated English
    # paragraph is not enough to open the record-image workflow.
    has_record_payload = bool(
        features.script_count >= 1
        or features.speaker_count >= 1
        or features.english_count >= 1
    )
    structural = bool(
        features.record_table_count >= 1
        or (record_instruction_count >= 2 and has_record_payload)
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=5 + record_instruction_count + features.record_table_count * 3 + features.script_count,
        signals=(
            "题型/小节标题: 听后记录并转述信息",
            f"记录表/答案表: {features.record_table_count}",
            f"记录说明与控制提示: {record_instruction_count}",
            f"听力正文载荷: {features.script_count + features.speaker_count + features.english_count}",
        ),
        details={"recording_table_count": features.record_table_count},
        question_count=features.table_question_count or features.question_count,
    )


def _detect_info_retelling(corpus: _Corpus, family):
    starts = []
    for position, (_, text, style) in enumerate(corpus.paragraphs):
        if not _is_heading_like(text, style):
            continue
        if not (_RETELLING_SECTION_RE.search(text) or "信息转述及询问" in text):
            continue
        # “第二节：信息转述” belongs to the new record/retelling family,
        # while “第一节 信息转述” and the top-level old title belong to the
        # legacy information-retelling family.
        if (
            _RETELLING_SECTION_RE.search(text)
            and re.search(r"第二节", text)
            and _record_context(corpus, position)
        ):
            continue
        if _RECORD_TITLE_RE.search(text):
            continue
        starts.append(position)
    starts = tuple(starts)
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    asking_count = sum(
        bool(_ASKING_SECTION_RE.search(text))
        for _, text, _ in corpus.paragraphs[start:end]
    )
    structural = bool(
        (
            features.script_count >= 1
            or features.speaker_count >= 1
            or features.english_count >= 1
        )
        and (asking_count >= 1 or features.answer_marker_count >= 1 or features.script_count >= 1)
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=4 + features.script_count * 3 + asking_count + features.question_count,
        signals=(
            "信息转述节/套卷标题",
            f"录音稿/英文正文: {features.script_count + features.speaker_count + features.english_count}",
            f"询问信息节: {asking_count}",
        ),
    )


def _detect_imitation(corpus: _Corpus, family):
    starts = _heading_positions(corpus, _IMITATION_HEADING_RE)
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    source_labels = sum(
        bool(_SOURCE_LABEL_RE.match(text))
        for _, text, _ in corpus.paragraphs[start:end]
    )
    boxed_questions = sum(
        bool(re.match(r"^\s*\d+\s*[.．、）)].*模仿朗读", text, re.IGNORECASE))
        for _, text, _ in corpus.paragraphs[start:end]
    )
    # The family heading itself contains “模仿朗读”, so it must not count as
    # an operation/control signal.  Otherwise an ordinary note such as
    # “模仿朗读\nThis is a note.” would pass the detector merely because the
    # title repeats the registered marker.
    imitation_control_positions = [
        position
        for position, (_, text, _) in enumerate(corpus.paragraphs[start:end], start)
        if position not in set(starts)
        and _IMITATION_CONTROL_RE.search(text)
    ]
    imitation_controls = len(imitation_control_positions)
    procedure_lines = sum(
        bool(re.search(r"听以下|请听|准备|开始录音|停止录音|听完后|朗读", text, re.IGNORECASE))
        for position, (_, text, _) in enumerate(corpus.paragraphs[start:end], start)
        if position not in set(starts)
    )
    has_source_layout = bool(
        source_labels >= 1
        or boxed_questions >= 1
        or features.english_table_count >= 1
    )
    structural = bool(
        (features.english_count >= 1 or features.english_table_count >= 1)
        and (
            has_source_layout
            or (imitation_controls >= 1 and procedure_lines >= 1)
        )
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=4 + features.english_count + features.english_table_count * 3 + source_labels * 2,
        signals=(
            "题型标题: 模仿朗读",
            f"英文朗读载荷: {features.english_count + features.english_table_count}",
            f"来源标签/编号题块: {source_labels + boxed_questions}",
            f"朗读控制提示(排除标题): {imitation_controls}",
            f"朗读流程语句: {procedure_lines}",
        ),
        question_count=(
            boxed_questions
            or source_labels
            or features.english_table_count
            or features.question_count
        ),
    )


def _find_header(headers: Iterable[Any], aliases: tuple[str, ...]) -> int | None:
    values = {
        index: str(value).strip()
        for index, value in enumerate(headers, start=1)
        if value is not None and str(value).strip()
    }
    for alias in aliases:
        for index, value in values.items():
            if value == alias:
                return index
    for alias in aliases:
        for index, value in values.items():
            if alias in value and "释义" not in value:
                return index
    return None


def _valid_vocabulary_workbook(source: Any) -> bool:
    try:
        import openpyxl
    except ImportError:
        return False
    workbook = None
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            header = []
            for row in worksheet.iter_rows(min_row=1, max_row=1, values_only=True):
                header = list(row)
                break
            if (
                _find_header(header, ("单词名称", "单词")) is not None
                and _find_header(header, ("例句",)) is not None
            ):
                return True
        return False
    except Exception:
        # A malformed or unsupported workbook is simply not evidence of the
        # vocabulary family.  The parser layer reports the detailed load
        # failure separately; detection must fail closed and remain safe for
        # batch imports.
        return False
    finally:
        if workbook is not None:
            workbook.close()


_PROFILE_DETECTORS = {
    "info_acquisition": _detect_info_acquisition,
    "listening_choice": _detect_selection,
    "listening_response": _detect_response,
    "text_reading": _detect_text_reading,
    "info_retelling": _detect_info_retelling,
    "listening_record_retelling": _detect_record_retelling,
    "imitation_reading": _detect_imitation,
}


_PAPER_FAMILY_CODES = frozenset({
    "info_acquisition",
    "listening_choice",
    "listening_response",
    "info_retelling",
    "listening_record_retelling",
    "imitation_reading",
})
_FULL_PAPER_TITLE_RE = re.compile(r"听说(?:考试|测试题)|听说新题型", re.IGNORECASE)
_SPECIAL_TITLE_RE = re.compile(r"专项", re.IGNORECASE)
_EXAM_SCORE_SIGNAL_RE = re.compile(
    r"(?:共\s*[0-9０-９零〇一二两三四五六七八九十百]+\s*(?:小\s*题|道\s*题|题)|"
    r"满分\s*[：:]?\s*[0-9０-９零〇一二两三四五六七八九十百]+\s*分|答题时间)",
    re.IGNORECASE,
)


def _paragraph_texts(value: Iterable[Any]) -> tuple[str, ...]:
    """Normalize paragraph-like values for document-level classification."""

    texts: list[str] = []
    for paragraph in value:
        if isinstance(paragraph, (tuple, list)) and len(paragraph) >= 2:
            text = paragraph[1]
        else:
            text = paragraph
        normalized = str(text or "").strip()
        if normalized:
            texts.append(normalized)
    return tuple(texts)


def classify_paper_category(
    paragraphs: Iterable[Any],
    *,
    evidences: Iterable[DocumentTypeEvidence] | None = None,
) -> PaperCategoryEvidence:
    """Classify a parsed Word document as题型专项 or听说考试.

    The decision is intentionally made after structural type detection.  A
    single major type is a专项 by default; a complete exam-bundle shape or a
    multi-section structure with independent title/score evidence is a套卷.
    Multiple types without those boundaries remain unresolved.  A visible
    ``专项`` label is only used as a conflict signal when it disagrees with a
    multi-section structure, so filenames and isolated words cannot silently
    change the result.
    """

    paragraph_values = tuple(paragraphs)
    text_values = _paragraph_texts(paragraph_values)
    detected = tuple(evidences) if evidences is not None else detect_document_types(
        paragraphs=paragraph_values,
    )
    paper_evidences = tuple(
        evidence for evidence in detected
        if evidence.code in _PAPER_FAMILY_CODES
    )
    detected_types = tuple(dict.fromkeys(
        evidence.doc_type for evidence in paper_evidences if evidence.doc_type
    ))
    if not paper_evidences:
        return PaperCategoryEvidence(
            paper_category=None,
            exam_form="unknown",
            status="not_applicable",
            confidence="none",
            signals=(),
        )

    full_text = "\n".join(text_values)
    bundle_shape = is_exam_paper_bundle(paragraph_values)
    full_title = bool(_FULL_PAPER_TITLE_RE.search(full_text))
    special_title = bool(_SPECIAL_TITLE_RE.search(full_text))
    score_signal_count = len(_EXAM_SCORE_SIGNAL_RE.findall(full_text))
    major_heading_count = sum(is_major_section_heading(text) for text in text_values)
    signals = [f"检测到 {len(detected_types)} 个独立大题型: {'、'.join(detected_types)}"]
    if bundle_shape:
        signals.append("正文同时满足完整听说套卷的题型/分值结构")
    if full_title:
        signals.append("正文标题包含听说考试/听说测试题")
    if special_title:
        signals.append("正文包含专项标记（仅作辅助证据）")
    signals.append(f"大题/小节边界: {major_heading_count}")
    signals.append(f"题量/分值/答题时间信号: {score_signal_count}")

    if len(detected_types) >= 2 and special_title and not bundle_shape and not full_title:
        return PaperCategoryEvidence(
            paper_category=None,
            exam_form="unknown",
            status="conflict",
            confidence="low",
            detected_types=detected_types,
            signals=tuple(signals + ["专项标记与多大题型结构冲突，等待确认"]),
        )

    has_full_paper_structure = bool(
        full_title
        or (major_heading_count >= 2 and score_signal_count >= 2)
    )
    if bundle_shape or (len(detected_types) >= 2 and has_full_paper_structure):
        confidence = "high" if bundle_shape else "medium"
        return PaperCategoryEvidence(
            paper_category="听说考试",
            exam_form="paper",
            status="suggested",
            confidence=confidence,
            detected_types=detected_types,
            signals=tuple(signals),
        )

    if len(detected_types) >= 2:
        return PaperCategoryEvidence(
            paper_category=None,
            exam_form="unknown",
            status="conflict",
            confidence="low",
            detected_types=detected_types,
            signals=tuple(signals + [
                "检测到多个题型，但缺少完整套卷边界/题量分值证据，等待确认",
            ]),
        )

    # A single detected type remains a专项 unless the body itself contains a
    # strong full-paper title and several independent exam signals.  This
    # handles truncated/legacy documents without letting one heading decide.
    if full_title and not special_title and major_heading_count >= 2 and score_signal_count >= 2:
        return PaperCategoryEvidence(
            paper_category="听说考试",
            exam_form="paper",
            status="suggested",
            confidence="medium",
            detected_types=detected_types,
            signals=tuple(signals + ["单一已识别题型，但正文保留完整套卷标题和结构信号"]),
        )

    return PaperCategoryEvidence(
        paper_category="题型专项",
        exam_form="special",
        status="suggested",
        confidence="high" if special_title else "medium",
        detected_types=detected_types,
        signals=tuple(signals),
    )


def _detect_generic(corpus: _Corpus, family):
    starts = _positions(
        corpus,
        lambda text, style: _is_heading_like(text, style)
        and _has_marker(family.content_markers, text),
    )
    if not starts:
        return None
    start, end = corpus.window(starts)
    features = _features(corpus, start, end)
    structural = bool(
        features.payload_count >= 1
        and (features.question_count >= 1 or features.script_count >= 1 or features.english_count >= 1)
    )
    if not structural:
        return None
    return _evidence(
        family,
        starts,
        features,
        structural=structural,
        score=2 + features.payload_count + features.question_count,
        signals=(
            "注册表内容标记命中标题",
            f"结构载荷: {features.payload_count}",
            f"编号题干: {features.question_count}",
        ),
    )


def detect_document_types(
    source: Any = None,
    *,
    paragraphs: Iterable[Any] | None = None,
    paragraph_metadata: Iterable[Any] | None = None,
    blocks: Iterable[Any] | None = None,
) -> tuple[DocumentTypeEvidence, ...]:
    """Detect document families from actual content and structural evidence.

    ``source`` may be a real ``.docx``/``.xlsx`` path, or may be omitted when
    ``paragraphs`` (and optional document ``blocks``) are already loaded.  A
    real path is loaded here only when the caller has not supplied the shared
    structures; the unified segmenter passes its one-load result to avoid a
    second read.

    A detector only returns evidence when the heading and the corresponding
    question/material structure are both present.  This keeps type routing
    deterministic and prevents a stray sentence containing a type name from
    opening a parser.
    """

    path = None
    if source is not None and isinstance(source, (str, os.PathLike, Path)):
        path = Path(os.fspath(source))

    if path is not None and path.suffix.lower() == ".xlsx":
        if paragraphs is not None:
            return ()
        if path.is_file() and _valid_vocabulary_workbook(path):
            family = FAMILY_REGISTRY.get("vocabulary")
            if family is None:
                return ()
            return (
                DocumentTypeEvidence(
                    code=family.code,
                    doc_type=family.display_name,
                    score=10,
                    structural=True,
                    signals=("工作簿存在单词/例句双表头",),
                    details={"format": "xlsx", "header_valid": True},
                ),
            )
        return ()

    if paragraphs is None:
        if path is None or not path.is_file():
            return ()
        try:
            from .text_utils import load_paragraphs

            loaded = load_paragraphs(
                path,
                include_metadata=True,
                include_blocks=True,
            )
            paragraphs = loaded[0]
            paragraph_metadata = loaded[1] if len(loaded) > 1 else None
            if blocks is None and len(loaded) > 2:
                blocks = loaded[2]
        except Exception:
            # Detection is a gate, not the place to expose package/zip parser
            # errors.  An unreadable or malformed source contributes no type
            # evidence and will be reported by the parser boundary if needed.
            return ()

    corpus = _build_corpus(paragraphs, blocks, paragraph_metadata)
    if not corpus.paragraphs:
        return ()

    evidences = []
    for family in FAMILY_REGISTRY.values():
        if family.code == "vocabulary":
            continue
        detector = _PROFILE_DETECTORS.get(family.code, _detect_generic)
        evidence = detector(corpus, family)
        if evidence is not None:
            evidences.append(evidence)

    # The parser consumes a mixed exam in source order.  Registry order remains
    # the tie-breaker for equal/unknown positions, so adding a family does not
    # make results nondeterministic.
    registry_order = {family.code: index for index, family in enumerate(FAMILY_REGISTRY.values())}
    return tuple(
        sorted(
            evidences,
            key=lambda evidence: (
                evidence.start_positions[0] if evidence.start_positions else len(corpus.paragraphs),
                registry_order.get(evidence.code, len(registry_order)),
            ),
        )
    )


def detect_document_type(
    source: Any = None,
    *,
    paragraphs: Iterable[Any] | None = None,
    paragraph_metadata: Iterable[Any] | None = None,
    blocks: Iterable[Any] | None = None,
) -> str | None:
    """Return one type only when structure detection is unambiguous."""

    evidences = detect_document_types(
        source,
        paragraphs=paragraphs,
        paragraph_metadata=paragraph_metadata,
        blocks=blocks,
    )
    if len(evidences) != 1:
        return None
    return evidences[0].doc_type


__all__ = [
    "DocumentTypeEvidence",
    "PaperCategoryEvidence",
    "classify_paper_category",
    "detect_document_type",
    "detect_document_types",
]
